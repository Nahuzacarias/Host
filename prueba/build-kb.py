#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de Base de Conocimiento para el Asistente Interno (Host iPPA)
Extrae información de Base de datos iPPA, Media/Host.xlsx, DOCX, PDF, XLS/XLSX y manuales.
"""

import os
import sys
import json
import shutil
import re
from pathlib import Path

# Setup encoding
sys.stdout.reconfigure(encoding='utf-8')

try:
    import pypdf
    import docx
    import openpyxl
    import xlrd
except ImportError as e:
    print(f"Error importando dependencias: {e}")
    sys.exit(1)

BASE_DIR = Path(__file__).resolve().parent.parent
PRUEBA_DOCS_DIR = BASE_DIR / "prueba" / "docs"
PRUEBA_DOCS_DIR.mkdir(parents=True, exist_ok=True)

def clean_text(text):
    if not text:
        return ""
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    text = re.sub(r'[ \t]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_pdf(path):
    try:
        reader = pypdf.PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ''
            if txt.strip():
                pages.append(f"--- Página {i+1} ---\n" + txt)
        return clean_text('\n\n'.join(pages))
    except Exception as e:
        print(f"Error leyendo PDF {path}: {e}")
        return ""

def extract_docx(path):
    try:
        doc = docx.Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    table_texts.append(" | ".join(cells))
        
        full_text = "\n".join(paragraphs)
        if table_texts:
            full_text += "\n\n--- Tablas y Datos ---\n" + "\n".join(table_texts)
        return clean_text(full_text)
    except Exception as e:
        print(f"Error leyendo DOCX {path}: {e}")
        return ""

def extract_xlsx(path):
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        sheets_out = []
        for sheetname in wb.sheetnames:
            sheet = wb[sheetname]
            rows_out = []
            for row in sheet.iter_rows(values_only=True):
                filtered = [str(c).strip() if c is not None else "" for c in row]
                if any(filtered):
                    rows_out.append(" | ".join(filtered))
            if rows_out:
                sheets_out.append(f"=== Hoja: {sheetname} ===\n" + "\n".join(rows_out))
        return clean_text('\n\n'.join(sheets_out))
    except Exception as e:
        try:
            book = xlrd.open_workbook(str(path))
            sheets_out = []
            for sheet in book.sheets():
                rows_out = []
                for row_idx in range(sheet.nrows):
                    row = sheet.row_values(row_idx)
                    filtered = [str(c).strip() if c is not None else "" for c in row]
                    if any(filtered):
                        rows_out.append(" | ".join(filtered))
                if rows_out:
                    sheets_out.append(f"=== Hoja: {sheet.name} ===\n" + "\n".join(rows_out))
            return clean_text('\n\n'.join(sheets_out))
        except Exception as e2:
            print(f"Error leyendo XLSX/XLS {path}: {e} / {e2}")
            return ""

def chunk_text(name, category, text, chunk_size=800, overlap=120):
    text = clean_text(text)
    if not text:
        return []
    
    chunks = []
    lines = text.split('\n')
    current_chunk = []
    current_len = 0
    
    for line in lines:
        line_len = len(line) + 1
        if current_len + line_len > chunk_size and current_chunk:
            chunk_content = '\n'.join(current_chunk).strip()
            if len(chunk_content) > 30:
                chunks.append({
                    "source": name,
                    "category": category,
                    "text": chunk_content
                })
            keep_lines = current_chunk[-2:] if len(current_chunk) >= 2 else current_chunk[-1:]
            current_chunk = [l for l in keep_lines]
            current_len = sum(len(l) + 1 for l in current_chunk)
            
        current_chunk.append(line)
        current_len += line_len
        
    if current_chunk:
        chunk_content = '\n'.join(current_chunk).strip()
        if len(chunk_content) > 30:
            chunks.append({
                "source": name,
                "category": category,
                "text": chunk_content
            })
            
    return chunks

def main():
    print("Iniciando extracción y compilación de Base de Conocimiento iPPA...")
    
    documents = []
    all_chunks = []
    processed_paths = set()
    
    sources_to_process = [
        {
            "path": BASE_DIR / "Media" / "Host.xlsx",
            "category": "Datos de Propiedades y Procedimientos Clave",
            "name": "Host.xlsx (Propiedades, Links, Claves y Checklists)"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Host" / "202512 Manual Host.pdf",
            "category": "Manual de Procedimientos",
            "name": "Manual Host 2025-2026"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Host" / "CheckTM.docx",
            "category": "Checklists de Turnos",
            "name": "Checklist Turno Mañana (TM)"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Host" / "CheckTT.docx",
            "category": "Checklists de Turnos",
            "name": "Checklist Turno Tarde (TT)"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Host" / "CheckTN.docx",
            "category": "Checklists de Turnos",
            "name": "Checklist Turno Noche (TN)"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Host" / "Recorrido experiencia iPPA.xlsx",
            "category": "Experiencia del Huésped",
            "name": "Recorrido Experiencia iPPA"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "HSK" / "Procedimiento reporte HSK.docx",
            "category": "Housekeeping",
            "name": "Procedimiento Reporte HSK"
        },
        {
            "path": BASE_DIR / "Guia_Aldea_Andina_Bariloche_servicios_gastronomia.docx",
            "category": "Guías de Destino",
            "name": "Guía Aldea Andina Bariloche (Servicios y Gastronomía)"
        },
        {
            "path": BASE_DIR / "Procedimiento facturacion AA.docx",
            "category": "Facturación y Cobranzas",
            "name": "Procedimiento Facturación Aldea Andina"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Aldea Andina" / "2. Operaciones" / "Aldea Andina by iPPA- Ficha Técnica.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Aldea Andina"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "BAH Congreso" / "2. Operaciones" / "Ficha Técnica - BAH Congreso.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - BAH Congreso"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "BAH Congreso" / "2. Operaciones" / "Info Accesos.xlsx",
            "category": "Accesos",
            "name": "Info Accesos - BAH Congreso"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Baires Belgrano" / "2. Operaciones" / "0. Info Propiedad" / "iP Hoteles - Ficha Técnica - Baires Belgrano by iPPA.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Baires Belgrano"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Filiberto" / "2. Operaciones" / "1.Manager" / "Guia de Operaciones - Filiberto.docx",
            "category": "Guías de Operaciones",
            "name": "Guía de Operaciones - Filiberto"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Filiberto" / "2. Operaciones" / "3.Host" / "Procedimientos" / "Procedimiento de cobranza WeTravel.docx",
            "category": "Facturación y Cobranzas",
            "name": "Procedimiento de Cobranza WeTravel"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Filiberto" / "2. Operaciones" / "6. Info Hotel" / "iP Hoteles - Ficha Técnica - Hotel Filiberto.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Hotel Filiberto"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Hache Palermo" / "2. Operaciones" / "0. Info Propiedad" / "Accesos Servicios y varios traspaso.docx",
            "category": "Accesos",
            "name": "Accesos Servicios - Hache Palermo"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Hache Palermo" / "2. Operaciones" / "0. Info Propiedad" / "Hache Palermo by iPPA - Ficha Técnica.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Hache Palermo"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Milenia Soho" / "2. Operaciones" / "6. Info Propiedad" / "iP Hoteles - Ficha Técnica - Milenia.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Milenia Soho"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Milenia Soho" / "4. Comercial" / "Política de Reservas Propietarios" / "iPPA Hoteles - Politica de uso propietarios.pdf",
            "category": "Políticas",
            "name": "Política de Uso Propietarios iPPA"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Milenia Vidt" / "1. Dirección" / "Milenia Vidt by iPPA - Ficha Técnica.xlsx",
            "category": "Fichas Técnicas",
            "name": "Ficha Técnica - Milenia Vidt"
        },
        {
            "path": BASE_DIR / "Base de datos iPPA" / "Clientes" / "Milenia Vidt" / "2. Operaciones" / "Proveedores.xlsx",
            "category": "Proveedores",
            "name": "Proveedores - Milenia Vidt"
        }
    ]
    
    # Record paths from base sources
    for item in sources_to_process:
        if "path" in item and item["path"]:
            processed_paths.add(str(item["path"].resolve()))
            processed_paths.add(item["path"].name.lower())

    # Dynamically scan for ANY additional files in prueba/docs/
    allowed_exts = [".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".txt", ".md"]
    skip_names = ["leeme.txt", "index.json", "knowledge-base.json"]
    
    for extra_file in PRUEBA_DOCS_DIR.iterdir():
        if not extra_file.is_file():
            continue
        if extra_file.name.lower() in skip_names or extra_file.name.lower() in processed_paths:
            continue
        if extra_file.suffix.lower() in allowed_exts:
            clean_name = extra_file.stem.replace('_', ' ').replace('-', ' ').title()
            sources_to_process.append({
                "path": extra_file,
                "category": "Documentos Adicionales",
                "name": clean_name
            })
            processed_paths.add(extra_file.name.lower())

    extracted_files = []
    
    for item in sources_to_process:
        path = item.get("path")
        name = item["name"]
        cat = item["category"]
        
        if "direct_content" in item:
            text = item["direct_content"]
        else:
            if not path or not path.exists():
                print(f"Advertencia: Archivo no encontrado: {path}")
                continue
            
            ext = path.suffix.lower()
            if ext == ".pdf":
                text = extract_pdf(path)
            elif ext in [".docx", ".doc"]:
                text = extract_docx(path)
            elif ext in [".xlsx", ".xls", ".csv"]:
                text = extract_xlsx(path)
            elif ext in [".txt", ".md"]:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
            else:
                continue

        if not text:
            print(f"Advertencia: No se pudo extraer texto de {name}")
            continue

        print(f"✓ Procesado: {name} ({len(text):,} caracteres)")
        
        doc_record = {
            "id": re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-'),
            "name": name,
            "category": cat,
            "char_count": len(text),
            "text": text
        }
        documents.append(doc_record)
        
        chunks = chunk_text(name, cat, text, chunk_size=900, overlap=120)
        all_chunks.extend(chunks)
        
        if path and path.exists() and not (PRUEBA_DOCS_DIR / path.name).exists():
            try:
                shutil.copy2(path, PRUEBA_DOCS_DIR / path.name)
            except Exception:
                pass
        extracted_files.append(path.name if path and path.exists() else name)

    kb_payload = {
        "generated_at": "2026-08-22",
        "total_documents": len(documents),
        "total_chunks": len(all_chunks),
        "categories": sorted(list(set(d["category"] for d in documents))),
        "documents": [
            {
                "id": d["id"],
                "name": d["name"],
                "category": d["category"],
                "char_count": d["char_count"]
            }
            for d in documents
        ],
        "chunks": all_chunks
    }
    
    kb_file = PRUEBA_DOCS_DIR / "knowledge-base.json"
    with open(kb_file, "w", encoding="utf-8") as f:
        json.dump(kb_payload, f, ensure_ascii=False, indent=2)
        
    print(f"\n¡Éxito! Base de conocimiento generada en {kb_file}")
    print(f"Total documentos: {len(documents)}")
    print(f"Total fragmentos (chunks): {len(all_chunks)}")
    
    index_files = [f.name for f in PRUEBA_DOCS_DIR.iterdir() if f.is_file() and f.suffix.lower() in [".pdf", ".docx", ".txt", ".md", ".xlsx", ".csv"] and f.name not in ["index.json", "knowledge-base.json"]]
    with open(PRUEBA_DOCS_DIR / "index.json", "w", encoding="utf-8") as f:
        json.dump({"files": sorted(index_files)}, f, ensure_ascii=False)
        
    print("index.json actualizado correctamente.")

if __name__ == "__main__":
    main()
